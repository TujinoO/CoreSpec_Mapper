from __future__ import annotations

"""Build the V5.3 Stable Markdown deliverables as styled, reviewable DOCX files."""

from pathlib import Path
from typing import Iterable
import argparse
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = (
    "CoreSpec_Mapper_V5_3_稳定版技术方法与系统设计_2026-07-21.md",
    "CoreSpec_Mapper_V5_3_稳定版用户使用指南_2026-07-21.md",
)

RELEASE_VERSION = "V5.3.0 Stable"
RELEASE_DATE = "2026-07-21"


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    value = OxmlElement("w:cantSplit")
    properties.append(value)


def _set_cell_no_wrap(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    if properties.find(qn("w:noWrap")) is None:
        properties.append(OxmlElement("w:noWrap"))


def _set_run_font(run, latin: str = "Aptos", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def _set_picture_alt(shape, description: str) -> None:
    properties = shape._inline.docPr
    properties.set("descr", description)
    properties.set("title", description)


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = f" {instruction} "
    run = paragraph.add_run()._r
    run.append(begin)
    run.append(instruction_node)
    run.append(separate)
    run.append(text)
    run.append(end)


def _plain_inline_text(text: str) -> str:
    """Return heading text without the small Markdown constructs we support."""

    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def _toc_entries(lines: list[str]) -> list[tuple[int, str, str]]:
    """Collect deterministic static-TOC entries from Markdown headings."""

    entries: list[tuple[int, str, str]] = []
    in_code = False
    for raw in lines:
        stripped = raw.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading is None:
            continue
        markdown_level = len(heading.group(1))
        word_level = min(max(markdown_level - 1, 1), 4)
        anchor = f"v5_heading_{len(entries) + 1:03d}"
        entries.append((word_level, _plain_inline_text(heading.group(2)), anchor))
    return entries


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph_xml = paragraph._p
    paragraph_xml.insert(1 if paragraph_xml.pPr is not None else 0, start)
    paragraph_xml.append(end)


def _add_internal_link(paragraph, text: str, anchor: str, color: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), "Aptos")
    run_fonts.set(qn("w:hAnsi"), "Aptos")
    run_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    properties.append(run_fonts)
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    properties.append(color_element)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    properties.append(underline)
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _new_numbering_instance(document: Document, start_at: int = 1) -> int:
    """Create a fresh List Number instance so each Markdown list restarts at 1."""

    numbering = document.part.numbering_part.element
    style_num_id = str(document.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(
        item for item in numbering.findall(qn("w:num"))
        if item.get(qn("w:numId")) == style_num_id
    )
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), str(start_at))
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    current = properties.find(qn("w:numPr"))
    if current is not None:
        properties.remove(current)
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.append(level)
    num_properties.append(number)
    properties.append(num_properties)


def _cover_title(title: str) -> str:
    replacements = (
        ("技术路线与方法总体设计", "技术路线与方法\n总体设计"),
        ("光谱数据库与识别能力说明", "光谱数据库与\n识别能力说明"),
        ("NC-1 实测验收与回归报告", "NC-1 实测验收与\n回归报告"),
        ("文档索引与版本说明", "文档索引与\n版本说明"),
        ("实施与验证报告", "实施与\n验证报告"),
        ("自适应矿物证据引擎软件总体设计方案", "自适应矿物证据引擎\n软件总体设计方案"),
        ("私有光谱数据库清单与识别能力", "私有光谱数据库清单与\n识别能力"),
        ("桌面版使用与 V4 迁移指南", "桌面版使用与 V4 迁移\n指南"),
        ("实测验收记录（", "实测验收记录\n（"),
    )
    for source, replacement in replacements:
        if source in title:
            return title.replace(source, replacement)
    return title


def _configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)

    for index, size, color, before, after in (
        (1, 16, "2E74B5", 18, 10),
        (2, 13, "2E74B5", 14, 7),
        (3, 12, "1F4D78", 10, 5),
        (4, 11, "1F4D78", 8, 4),
    ):
        style = document.styles[f"Heading {index}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in document.styles:
        style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Cascadia Mono"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(8.5)
        style.paragraph_format.left_indent = Cm(0.45)
        style.paragraph_format.right_indent = Cm(0.2)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(6)

    if "Document Note" not in document.styles:
        style = document.styles.add_style("Document Note", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor.from_string("566B7C")
        style.paragraph_format.left_indent = Cm(0.35)
        style.paragraph_format.space_after = Pt(6)

    for level, size, indent, color in (
        (1, 7.8, 0.0, "153B5C"),
        (2, 7.4, 0.50, "2C6D90"),
        (3, 7.0, 1.00, "526273"),
        (4, 6.8, 1.45, "607487"),
    ):
        style_name = f"Static TOC {level}"
        if style_name not in document.styles:
            style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.left_indent = Cm(indent)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0.5)
        style.paragraph_format.keep_with_next = False

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("CoreSpec Mapper V5.3 Stable  ·  发布文档")
    _set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("718399")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("CoreSpec Mapper V5.3 Stable   |   ")
    _set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("718399")
    _field(footer, "PAGE", "1")

    document.core_properties.title = title
    document.core_properties.subject = "CoreSpec Mapper V5.3 Stable technical method and user documentation"
    document.core_properties.author = "CoreSpec Mapper V5.3 Project"
    document.core_properties.keywords = "hyperspectral, drill core, mineral mapping, V5.3, NC-1, stripe suppression"


def _add_cover(document: Document, title: str, toc_entries: list[tuple[int, str, str]]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "src" / "corespec_mapper" / "resources" / "corespec_logo.png"
    if logo.is_file():
        shape = paragraph.add_run().add_picture(str(logo), width=Cm(2.2))
        _set_picture_alt(shape, "CoreSpec Mapper 应用标志")

    brand = document.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.paragraph_format.space_before = Pt(14)
    run = brand.add_run("CORESPEC MAPPER  ·  STABLE RELEASE")
    _set_run_font(run, "Aptos Display")
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("2684C7")

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(42)
    heading.paragraph_format.space_after = Pt(18)
    run = heading.add_run(_cover_title(title))
    _set_run_font(run, "Aptos Display")
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("153B5C")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(22)
    run = meta.add_run(
        f"{RELEASE_VERSION}  ·  {RELEASE_DATE}\n"
        "九步桌面 · 自适应三档识别 · V14 重复列与主导走廊终检"
    )
    _set_run_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("526273")

    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(72)
    run = notice.add_run("工程质量等级不等同于矿物学准确率；无像元级真值时不得作精度宣称。")
    _set_run_font(run)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string("7A4E00")
    document.add_page_break()

    toc_title = document.add_paragraph("目录", style="Heading 1")
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_note = document.add_paragraph(style="Document Note")
    toc_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = toc_note.add_run("静态章节目录 · 可点击跳转 · 不依赖 Word 域更新")
    _set_run_font(note_run)
    for level, text, anchor in toc_entries:
        paragraph = document.add_paragraph(style=f"Static TOC {level}")
        _add_internal_link(paragraph, text, anchor, "153B5C" if level == 1 else "2C6D90")
    document.add_page_break()


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")


def _add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            _set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run)
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Cascadia Mono")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("8A3E16")
        else:
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label}（{url}）")
            _set_run_font(run)
            run.font.color.rgb = RGBColor.from_string("1769AA")
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_run_font(run)


def _table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        row = [item.strip() for item in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", item.replace(" ", "")) for item in rows[1]):
        rows.pop(1)
    return rows, index


def _display_width(value: str) -> int:
    """Approximate visual width, counting CJK glyphs wider than Latin glyphs."""

    plain = _plain_inline_text(value)
    return sum(2 if ord(character) > 0x7F else 1 for character in plain)


def _column_widths(rows: list[list[str]], columns: int, total_cm: float = 16.51) -> list[float]:
    """Return deterministic fixed table widths that fit the compact page geometry."""

    scores = []
    for column in range(columns):
        score = max(
            (_display_width(row[column]) for row in rows if column < len(row)),
            default=8,
        )
        scores.append(float(min(max(score, 6), 42)))
    minimum = 1.15 if columns <= 5 else 0.78
    if minimum * columns >= total_cm:
        return [total_cm / columns] * columns
    remaining = total_cm - minimum * columns
    total_score = sum(scores) or float(columns)
    return [minimum + remaining * score / total_score for score in scores]


def _set_fixed_table_layout(table) -> None:
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def _apply_exact_table_geometry(table, widths_cm: list[float]) -> None:
    """Apply compact_reference_guide fixed 9360-DXA table geometry."""

    widths = [round(value / 2.54 * 1440) for value in widths_cm]
    widths[-1] += 9360 - sum(widths)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), "9360")
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    margins = properties.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        properties.append(margins)
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:type"), "dxa")
        element.set(qn("w:w"), str(value))
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(widths[index]))


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_fixed_table_layout(table)
    widths = _column_widths(rows, columns)
    _apply_exact_table_geometry(table, widths)
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Cm(width)
    longest_token = max(
        (
            len(token)
            for row in rows
            for value in row
            for token in re.findall(r"[A-Za-z0-9_.:/\\-]+", value)
        ),
        default=0,
    )
    table_font_size = 7.0 if columns >= 8 or longest_token > 28 else 8.5
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        _prevent_row_split(row)
        if row_index == 0:
            _set_repeat_header(row)
        for column, cell in enumerate(row.cells):
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                _set_cell_shading(cell, "DCEAF5")
            elif row_index % 2 == 0:
                _set_cell_shading(cell, "F5F8FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if row_index == 0:
                paragraph.paragraph_format.keep_with_next = True
            cell_value = values[column] if column < len(values) else ""
            compact_value = cell_value.replace(" ", "")
            if columns >= 8 and re.fullmatch(r"[0-9][0-9.,%×→/–—-]*", compact_value):
                _set_cell_no_wrap(cell)
            _add_inline(paragraph, cell_value)
            for run in paragraph.runs:
                run.font.size = Pt(table_font_size)
                if row_index == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _render_markdown(
    document: Document,
    lines: list[str],
    toc_entries: list[tuple[int, str, str]],
) -> None:
    index = 0
    heading_index = 0
    active_numbering_id: int | None = None
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if stripped.startswith("```"):
            active_numbering_id = None
            if not in_code:
                in_code = True
                code_language = stripped[3:].strip()
                code_lines = []
            else:
                if code_language.casefold() == "mermaid":
                    note = document.add_paragraph("流程图定义（Mermaid 源码）", style="Document Note")
                    note.paragraph_format.keep_with_next = True
                paragraph = document.add_paragraph(style="Code Block")
                _set_cell_shading_like_paragraph(paragraph, "F0F3F5")
                run = paragraph.add_run("\n".join(code_lines))
                _set_run_font(run, "Cascadia Mono")
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped:
            active_numbering_id = None
            index += 1
            continue
        if stripped.startswith("|"):
            active_numbering_id = None
            rows, index = _table_rows(lines, index)
            _add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            active_numbering_id = None
            markdown_level = len(heading.group(1))
            word_level = min(max(markdown_level - 1, 1), 4)
            paragraph = document.add_paragraph(style=f"Heading {word_level}")
            _add_inline(paragraph, heading.group(2))
            if heading_index < len(toc_entries):
                _add_bookmark(paragraph, toc_entries[heading_index][2], heading_index + 1)
            heading_index += 1
            index += 1
            continue
        if re.fullmatch(r"-{3,}", stripped):
            active_numbering_id = None
            index += 1
            continue
        bullet = re.match(r"^(\s*)[-*+]\s+(.+)$", raw)
        numbered = re.match(r"^(\s*)(\d+)[.)]\s+(.+)$", raw)
        if bullet or numbered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            indent = bullet.group(1) if bullet else numbered.group(1)
            item_text = bullet.group(2) if bullet else numbered.group(3)
            paragraph.paragraph_format.left_indent = Inches(0.375 + min(len(indent) // 2, 3) * 0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.188)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            if bullet:
                active_numbering_id = None
            else:
                if active_numbering_id is None:
                    active_numbering_id = _new_numbering_instance(document, int(numbered.group(2)))
                _apply_numbering(paragraph, active_numbering_id)
            _add_inline(paragraph, item_text)
            index += 1
            continue
        if stripped.startswith(">"):
            active_numbering_id = None
            paragraph = document.add_paragraph(style="Document Note")
            _add_inline(paragraph, stripped.lstrip("> "))
            _set_cell_shading_like_paragraph(paragraph, "EEF5FA")
            index += 1
            continue

        active_numbering_id = None
        parts = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "|", "```", ">"))
                or re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)", candidate)
            ):
                break
            parts.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        paragraph_text = " ".join(parts)
        _add_inline(paragraph, paragraph_text)
        following = index
        while following < len(lines) and not lines[following].strip():
            following += 1
        if following < len(lines) and paragraph_text.rstrip().endswith((":", "：")):
            next_line = lines[following].strip()
            if (
                next_line.startswith(("|", "```", "- ", "* ", "+ "))
                or re.match(r"^\d+[.)]\s+", next_line)
            ):
                paragraph.paragraph_format.keep_with_next = True


def _set_cell_shading_like_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _append_validation_figures(document: Document) -> None:
    nc1_root = (
        ROOT
        / "artifacts"
        / "NC1_V5_3_Adaptive_Regression"
        / "NC1_V5_3_Adaptive_Regression"
        / "project_calibrated_v14_dominant_corridor_20260721"
        / "previews"
    )
    nc1_figures = (
        (
            nc1_root / "comparison_balanced.png",
            "图 A-1  NC-1 V14 平衡档：默认地质解释基线",
        ),
        (
            nc1_root / "comparison_sensitive.png",
            "图 A-2  NC-1 V14 敏感档：找漏与边界检查",
        ),
    )
    available_nc1 = [(path, caption) for path, caption in nc1_figures if path.is_file()]
    if available_nc1:
        document.add_page_break()
        document.add_paragraph("附录 A  NC-1 V14 全景分类复核图", style="Heading 1")
        note = document.add_paragraph(style="Document Note")
        _add_inline(note, "图像为 5,446 行全深度压缩概览，用于比较三档数量差异、空间连续性和重复列风险；像元级结论以 ENVI 栅格、质量报告和 run_manifest.json 为准。")
        for path, caption in available_nc1:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            shape = paragraph.add_run().add_picture(str(path), height=Cm(17.8))
            _set_picture_alt(shape, caption)
            caption_paragraph = document.add_paragraph()
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption_paragraph.add_run(caption)
            _set_run_font(run)
            run.font.size = Pt(8.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor.from_string("526273")
        return

    figures = (
        (
            ROOT / "validation_runs" / "V5_Final_Validation" / "3dssz_lowres_v5_final2" / "previews" / "comparison_balanced.png",
            "图 A-1  3DSSZ 连续低分辨率立方体：平衡档五个光谱竞争族对比（第一列为 1600 nm 背景）",
        ),
        (
            ROOT / "validation_runs" / "V5_Final_Validation" / "zkh3_lowres_v5_final2" / "previews" / "comparison_balanced.png",
            "图 A-2  ZKH3 连续低分辨率立方体：平衡档五个光谱竞争族对比（第一列为 1600 nm 背景）",
        ),
        (
            ROOT / "validation_runs" / "V5_Final_Validation" / "zkh3_lowres_v5_final2" / "previews" / "aloh_wavelength_subtype.png",
            "图 A-3  ZKH3 白云母—伊利石族 Al-OH 短/中/长波亚型预览；该亚型是峰位描述，不等同于确定矿物相分离",
        ),
    )
    available = [(path, caption) for path, caption in figures if path.is_file()]
    if not available:
        return
    document.add_page_break()
    document.add_paragraph("附录 A  可视化复核图", style="Heading 1")
    for path, caption in available:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        width = Inches(4.0 if path.name == "aloh_wavelength_subtype.png" else 6.35)
        paragraph.add_run().add_picture(str(path), width=width)
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.keep_with_next = False
        run = caption_paragraph.add_run(caption)
        _set_run_font(run)
        run.font.size = Pt(8.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string("526273")


def build(source: Path, destination: Path) -> Path:
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = source.stem
    if lines and lines[0].startswith("# "):
        title = lines[0][2:].strip()
        lines = lines[1:]
    toc_entries = _toc_entries(lines)
    document = Document()
    _configure_document(document, title)
    _add_cover(document, title, toc_entries)
    _render_markdown(document, lines, toc_entries)
    if "稳定版技术方法与系统设计" in source.name:
        _append_validation_figures(document)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", action="append", help="Markdown source; repeatable")
    parser.add_argument("--output-dir", default=str(ROOT / "docs"))
    args = parser.parse_args(list(argv) if argv is not None else None)
    sources = [Path(value) for value in args.source] if args.source else [ROOT / "docs" / name for name in DOCS]
    output = Path(args.output_dir)
    for source in sources:
        destination = output / f"{source.stem}.docx"
        print(build(source, destination))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
